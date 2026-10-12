"""Design system: style resolution and Word font embedding."""

import zipfile

import docx
import pytest

from services import design_system as ds


@pytest.mark.parametrize(
    "name, expected",
    [
        ("Weekly Update.docx", ds.BRAND),
        ("Investor briefing Q4", ds.BRAND),
        ("Tech spec - pump head", ds.BRAND),
        ("Memo to Hatch", ds.BRAND),
        ("CAPA-004.docx", ds.CONTROLLED),
        ("SOP-011 Supplier Control", ds.CONTROLLED),
        ("Risk Management Plan", ds.CONTROLLED),
        ("ISO 14971 report", ds.CONTROLLED),
        ("510(k) summary", ds.CONTROLLED),
        ("Quality Manual", ds.CONTROLLED),
        ("Design Verification Protocol DV-05", ds.CONTROLLED),
        ("Regulatory strategy", ds.CONTROLLED),
    ],
)
def test_resolve_doc_style_detects_controlled_documents(name, expected):
    assert ds.resolve_doc_style(None, name) == expected


def test_explicit_style_wins_over_detection():
    assert ds.resolve_doc_style("brand", "CAPA-004") == ds.BRAND
    assert ds.resolve_doc_style("controlled", "Weekly update") == ds.CONTROLLED
    assert ds.resolve_doc_style("nonsense", "Weekly update") == ds.BRAND


def test_body_words_do_not_trigger_controlled():
    # Only names/titles are classified; a weekly update mentioning risk is brand.
    assert ds.resolve_doc_style(None, "Weekly update", "Programme status") == ds.BRAND


def test_obfuscation_round_trips():
    key = "{0F7A2B3C-4D5E-4F60-8A9B-0C1D2E3F4A5B}"
    data = bytes(range(64))
    once = ds._obfuscate(data, key)
    assert once != data and once[32:] == data[32:]
    assert ds._obfuscate(once, key) == data


def test_embed_fonts_writes_parts_and_settings(tmp_path):
    if not (ds.FONTS_DIR / "Archivo-Regular.ttf").is_file():
        pytest.skip("font assets not present")
    document = docx.Document()
    document.add_paragraph("x")
    embedded = ds.embed_fonts(document)
    assert embedded == ["Archivo", "JetBrains Mono"]
    # Idempotent: a second call adds nothing new.
    assert ds.embed_fonts(document) == ["Archivo", "JetBrains Mono"]
    out = tmp_path / "f.docx"
    document.save(str(out))
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        font_table = z.read("word/fontTable.xml").decode()
        settings = z.read("word/settings.xml").decode()
        content_types = z.read("[Content_Types].xml").decode()
    assert sorted(n for n in names if n.startswith("word/fonts/")) == [
        f"word/fonts/font{i}.odttf" for i in range(1, 5)
    ]
    assert font_table.count("w:embedRegular") == 2 and font_table.count("w:embedBold") == 2
    assert 'w:name="Archivo"' in font_table and 'w:name="JetBrains Mono"' in font_table
    assert "w:embedTrueTypeFonts" in settings
    assert "obfuscatedFont" in content_types
    docx.Document(str(out))  # still a valid package
